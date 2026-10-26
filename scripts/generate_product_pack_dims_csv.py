#!/usr/bin/env python3
"""Generate suggested pack dimensions CSV from exports/product_weights_prod.csv.

Run locally (no Odoo):
  python scripts/generate_product_pack_dims_csv.py

Output: exports/product_pack_dims_suggested.csv
Edit columns length_cm,width_cm,height_cm then apply on prod via
apply_product_pack_dims_csv_shell.py
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WEIGHTS = ROOT / "exports" / "product_weights_prod.csv"
OUT_CSV = ROOT / "exports" / "product_pack_dims_suggested.csv"
OUT_XLSX = ROOT / "exports" / "product_pack_dims_suggested.xlsx"


def text_blob(row):
    parts = [
        row.get("product_name") or "",
        row.get("default_code") or "",
        row.get("internal_category") or "",
        row.get("shop_categories") or "",
    ]
    return " ".join(parts).lower()


def suggest_pack_dims_cm(row):
    blob = text_blob(row)
    shop = (row.get("shop_categories") or "").lower()
    categ = (row.get("internal_category") or "").lower()

    def dims(l, w, h):
        return (l, w, h)

    if re.search(r"\b(dual|2 x|two)\b.*\b(monitor|screen|display)\b", blob) or re.search(
        r"\b(monitor|screen|display).*\bbundle\b", blob
    ):
        if re.search(r"\b24[\"']|24 inch\b", blob):
            return dims(60, 40, 25)
        return dims(55, 38, 22)
    if (
        "cables & adapters" in shop
        or categ == "chargers / power adapters"
        or re.search(r"\b(cable|adapter|charger|hdmi|displayport|usb|ethernet extension)\b", blob)
    ):
        return dims(18, 12, 4)
    if "laptops" in shop or (
        categ == "computer systems"
        and re.search(r"\b(thinkpad|latitude|toughbook|cf-54|elitebook|probook)\b", blob)
    ):
        if re.search(r"\b(toughbook|cf-54)\b", blob):
            return dims(42, 32, 8)
        if re.search(r"\b(t15|p15|17[\"']|17 inch)\b", blob):
            return dims(40, 30, 6)
        return dims(38, 28, 5)
    if re.search(r"\b(optiplex|elite desk|9020|800 g2|tiny|m910)\b", blob):
        if re.search(r"\b(mini|tiny|m910|nuc)\b", blob):
            return dims(35, 30, 15)
        return dims(40, 35, 18)
    if re.search(r"\b(thinkstation|p510|workstation)\b", blob):
        return dims(50, 45, 30)
    if "monitors" in shop or "monitor" in blob:
        if re.search(r"\b32[\"']", blob):
            return dims(75, 45, 15)
        if re.search(r"\b27[\"']", blob):
            return dims(65, 42, 14)
        return dims(55, 35, 12)
    if "docks" in shop or re.search(r"\bdock\b", blob):
        return dims(28, 20, 8)
    if "projector" in blob:
        return dims(38, 30, 12)
    if "digital camera" in shop or re.search(r"\b(camer|powershot|lumix)\b", blob):
        return dims(15, 12, 8)
    if re.search(r"\b(jabra|headphone)\b", blob):
        return dims(22, 18, 10)
    if re.search(r"\b(mouse|keyboard)\b", blob):
        return dims(35, 15, 5)
    if re.search(r"\b(3750|48p)\b", blob):
        return dims(55, 40, 15)
    if re.search(r"\b(router|switch|cisco|meraki|access point)\b", blob):
        return dims(35, 25, 10)
    if re.search(r"\b(backpack|bag)\b", blob):
        return dims(45, 35, 15)
    if re.search(r"\b(speaker|conference cam)\b", blob):
        return dims(30, 25, 20)
    return dims(30, 25, 10)


def main():
    if not WEIGHTS.is_file():
        raise SystemExit(f"Missing {WEIGHTS}")

    with WEIGHTS.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "product_id",
        "default_code",
        "product_name",
        "weight_kg",
        "length_cm",
        "width_cm",
        "height_cm",
        "notes",
    ]
    out_rows = []
    for row in rows:
        l, wd, h = suggest_pack_dims_cm(row)
        out_rows.append(
            {
                "product_id": row["product_id"],
                "default_code": row.get("default_code") or "",
                "product_name": row.get("product_name") or "",
                "weight_kg": row.get("weight_kg") or "",
                "length_cm": l,
                "width_cm": wd,
                "height_cm": h,
                "notes": "Edit L/W/H then re-import via apply_product_pack_dims_csv_shell.py",
            }
        )

    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(out_rows)
    print(f"Wrote {len(out_rows)} rows -> {OUT_CSV}")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl not installed — skip XLSX (pip install openpyxl)")
        return

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "Re-Ware pack dimensions (cm) for AusPost / CartonPilot"
    readme["A1"].font = Font(bold=True)
    readme["A3"] = "Columns on sheet 'products': product_id must match Odoo product.template id."
    readme["A4"] = "Edit length_cm, width_cm, height_cm (bare product size is OK)."
    readme["A5"] = "Odoo adds foam allowance at quote time: Settings → Pack padding (cm) + extra kg."
    readme["A6"] = "After edits: save as CSV (product_id,length_cm,width_cm,height_cm) and apply on prod /tmp/"

    ws = wb.create_sheet("products")
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for rec in out_rows:
        ws.append([rec[k] for k in fieldnames])
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    print(f"Wrote Excel -> {OUT_XLSX}")


if __name__ == "__main__":
    main()
